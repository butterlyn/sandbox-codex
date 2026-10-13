"""Interactive reconstructions of image-backed appendix figures 27 and 28.

All plotted observations come from the register's saved worksheet values. The
reliability-standard shading in Figure 27 is explicitly derived; Figure 28's
Essential System Services area sums the four recorded reserve series.
"""
from __future__ import annotations

import datetime as dt
import json
import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "data/sources/2026-wem-esoo-data-register.xlsx"
PURPLE = "#702C82"
COLORS = {
    "Coal": "#444444", "Wind": "#9ED777", "Solar": "#FFB719",
    "Distillate": "#EF6657", "Gas": "#3DBDAA", "Waste": "#DFE689",
    "ESR (discharge)": "#AD56A7", "Demand-side participation": "#666BB6",
    "Unserved Energy": "#A32046", "Distributed PV": "#FBE287",
    "Essential System Services": "#DFF2F7", "Thermal headroom": "#F2DDF3",
    "Forced outages": "#C7C7C7", "Regulation Raise": "#CDCCE7",
    "Contingency Reserve Raise": "#DFF2F7", "Planned Outage": "#B9B9B9",
    "Forced Outage": "#D8D8D8", "Headroom": "#F2DDF3",
}


def layout(y_title: str) -> dict:
    return {
        "font": {"family": "Arial, sans-serif", "size": 12, "color": "#283C49"},
        "paper_bgcolor": "#FFFFFF", "plot_bgcolor": "#FFFFFF", "hovermode": "x unified",
        "dragmode": "zoom", "margin": {"l": 80, "r": 70, "t": 28, "b": 150},
        "legend": {"orientation": "h", "y": -0.28, "x": 0, "font": {"size": 11}, "traceorder": "normal"},
        "xaxis": {"title": {"text": "Date and time"}, "type": "date", "showgrid": False,
                  "tickformat": "%d %b<br>%H:%M", "dtick": 21600000, "automargin": True},
        "yaxis": {"title": {"text": y_title}, "showgrid": True, "gridcolor": "#E5EAF0",
                  "zerolinecolor": "#8796A4", "tickformat": ",.0f", "automargin": True},
    }


def base_spec(identifier: str, sheet: str, title: str, image: int) -> dict:
    return {
        "id": identifier, "sheet": sheet, "title": title, "section": "Appendices",
        "kind": "interactive", "source_format": "worksheet data behind image",
        "source_file": BOOK.name, "image_file": f"data/analysis/images/image{image}.png",
        "data": [], "layout": {}, "notes": [], "sources": [],
    }


def bubble_spec(workbook) -> dict:
    sheet = workbook["A4_F.27"]
    observations = [tuple(sheet.cell(row, col).value for col in (2, 3, 4)) for row in range(25, 325)]
    assert len(observations) == 300
    assert math.isclose(sum(row[2] for row in observations), 1.0, abs_tol=1e-12)
    assert sum(row[2] > 0 for row in observations) == 39
    spec = base_spec("supplement-a4-f27", sheet.title, sheet["A1"].value, 35)
    # Keep all 299 subsequent source rows, including the 261 zero-frequency
    # placeholders. Their zero marker sizes preserve the published appearance.
    remainder = observations[1:]
    frequencies = [row[2] for row in remainder]
    spec["data"] = [
        {
            "type": "scatter", "mode": "markers", "name": "Unserved-energy events",
            "x": [row[0] for row in remainder], "y": [row[1] for row in remainder],
            "marker": {"color": PURPLE, "opacity": 0.95, "size": frequencies,
                       "sizemode": "area", "sizeref": 2 * max(frequencies) / (43 ** 2),
                       "line": {"width": 0.5, "color": "#FFFFFF"}},
            "customdata": [[row[2], row[2] * 100, row_number] for row_number, row in enumerate(remainder, 26)],
            "meta": {"sourceValueCount": 299, "frequencyColumn": "D", "sourceRows": [26, 324]},
            "hovertemplate": "Average event duration: %{x:.1f} hours<br>Demand magnitude: %{y:.1f}%<br>Simulation frequency: %{customdata[1]:.4f}%<br>Source row: %{customdata[2]}<extra></extra>",
        },
        {
            "type": "scatter", "mode": "markers+text", "name": "No unserved energy",
            "x": [observations[0][0]], "y": [observations[0][1]],
            "marker": {"size": 64, "color": PURPLE, "opacity": 0.95, "line": {"width": 1, "color": "#FFFFFF"}},
            "text": [f"{observations[0][2] * 100:.0f}%"], "textposition": "middle center",
            "textfont": {"color": "#FFFFFF", "size": 17}, "cliponaxis": False,
            "customdata": [[observations[0][2], observations[0][2] * 100, 25]],
            "meta": {"sourceValueCount": 1, "frequencyColumn": "D", "sourceRows": [25, 25], "bubbleSizeNotToScale": True},
            "hovertemplate": "No unserved-energy event<br>Simulation frequency: %{customdata[1]:.4f}%<br>Source row: 25<extra></extra>",
        },
    ]
    # The WEM standard allows 0.0002% annual USE; 8,760 hours x 0.0002%
    # gives a 1.752 %-hour boundary when magnitude is a percentage of demand.
    threshold_percent_hours = 8760 * 0.0002
    start_x = threshold_percent_hours / 30
    curve_x = [start_x + (6 - start_x) * i / 250 for i in range(251)]
    path = f"M {start_x},30 L 6,30 " + " ".join(
        f"L {x:.8f},{threshold_percent_hours / x:.8f}" for x in reversed(curve_x)
    ) + " Z"
    spec["layout"] = layout("Magnitude (% of average unscheduled operational demand)")
    spec["layout"].update({
        "hovermode": "closest", "margin": {"l": 90, "r": 35, "t": 35, "b": 115},
        "xaxis": {"title": {"text": "Average event duration in simulation (hours)"}, "type": "linear",
                  "range": [0, 6], "tickmode": "array", "tickvals": [0, 1, 2, 3, 4, 5, 6],
                  "ticktext": ["0", "1", "2", "3", "4", "5", "≥ 6"], "automargin": True},
        "yaxis": {"title": {"text": "Magnitude (% of average unscheduled demand)"}, "range": [0, 30],
                  "tickmode": "array", "tickvals": [0, 10, 20, 30], "ticksuffix": "%", "automargin": True},
        "shapes": [{"type": "path", "path": path, "xref": "x", "yref": "y", "layer": "below",
                    "fillcolor": "rgba(120,120,120,0.30)", "line": {"color": "#888888", "width": 1.2}}],
        "annotations": [{"xref": "paper", "yref": "paper", "x": 0.98, "y": 0.99,
                         "xanchor": "right", "yanchor": "top", "showarrow": False,
                         "text": "Bubble area: simulation frequency<br>Shading: above annual reliability standard",
                         "font": {"size": 11, "color": "#555555"}, "bgcolor": "rgba(255,255,255,0.9)"}],
    })
    spec["sources"] = [
        {"sheet": sheet.title, "range": "B25:B324", "role": "Average event duration (hours)"},
        {"sheet": sheet.title, "range": "C25:C324", "role": "Magnitude (% of average demand)"},
        {"sheet": sheet.title, "range": "D25:D324", "role": "Simulation frequency (fraction; multiplied by 100 for percentage hover)"},
    ]
    spec["notes"] = [
        "All 300 worksheet rows are retained: 39 have positive frequency, including the no-event case; 261 are zero-frequency placeholders. Frequencies sum to 100%.",
        "The zero-point bubble represents 71.1875% of simulations and is not to scale, matching the source figure. Other bubble areas represent frequency.",
        "The grey reliability boundary is reconstructed from the 0.0002% annual unserved-energy standard and 8,760 hours: magnitude (%) × duration (hours) = 1.752. It is a derived reference curve, not an extra worksheet observation.",
    ]
    return spec


def temporal_specs(workbook) -> tuple[list[dict], dict]:
    sheet = workbook["A4_F.28"]
    rows = list(sheet.iter_rows(min_row=44, max_row=140, values_only=True))
    assert len(rows) == 97
    times = []
    for row in rows:
        date, time = row[1], row[2]
        time = time.time() if isinstance(time, dt.datetime) else time
        times.append(dt.datetime.combine(date.date(), time).isoformat())
    assert all((dt.datetime.fromisoformat(b) - dt.datetime.fromisoformat(a)).total_seconds() == 1800
               for a, b in zip(times, times[1:]))

    def values(column: str) -> list[float]:
        return [float(row[column_index_from_string(column) - 1]) for row in rows]

    def source(column: str, role: str | None = None) -> dict:
        return {"sheet": sheet.title, "range": f"{column}44:{column}140",
                "role": role or str(sheet[f"{column}43"].value)}

    def trace(column: str, name: str | None = None, *, stack: str | None = None,
              color: str | None = None, negate: bool = False, dash: str = "solid",
              secondary: bool = False, fill: bool = False) -> dict:
        label = name or str(sheet[f"{column}43"].value)
        raw = values(column)
        result = {
            "type": "scatter", "mode": "lines", "name": label, "x": times,
            "y": [-v for v in raw] if negate else raw,
            "line": {"color": color or COLORS.get(label, PURPLE), "width": 0.5 if stack else 2.2, "dash": dash},
            "customdata": raw, "connectgaps": False,
            "meta": {"sourceColumn": column, "sourceValueCount": len(raw), "sourceRows": [44, 140],
                     "yScaleFactor": 1, "yTransform": "negate" if negate else "identity"},
            "hovertemplate": "%{x|%d %b %Y %H:%M}<br>%{fullData.name}: %{y:,.2f}" + ("%" if secondary else " MW") + "<extra></extra>",
        }
        if stack:
            result.update({"stackgroup": stack, "fillcolor": color or COLORS.get(label, PURPLE), "stackgaps": "infer zero"})
        if fill:
            result.update({"fill": "tozeroy", "fillcolor": "rgba(174,90,169,0.58)"})
        if secondary:
            result["yaxis"] = "y2"
        return result

    common_notes = [
        "All 97 half-hourly samples are retained, from 1 July 2032 00:00 through 3 July 2032 00:00, using the worksheet's sample dates. The report describes this as a 2030–31 modelled event using the 2009–10 weather reference year.",
        "ESR charging (withdrawal) is plotted below zero; the original positive charging values remain in the trace source data. State of charge uses the right-hand percentage axis.",
    ]
    base_sources = [source("B", "Sample date"), source("C", "Sample time of day")]
    spec = base_spec("supplement-a4-f28-dispatch", sheet.title,
                     "Figure 28 — Modelled supply and demand during a two-day unserved-energy event", 36)
    spec["subtitle"] = "Top panel · generation, demand and storage state of charge"
    positive_columns = list("DEFGHIJKLM")
    spec["data"] = [trace(col, stack="dispatch") for col in positive_columns]
    reserve_columns = ["U", "V", "Z", "AA"]
    ess = [sum(v) for v in zip(*(values(col) for col in reserve_columns))]
    spec["data"].append({
        "type": "scatter", "mode": "lines", "name": "Essential System Services", "x": times, "y": ess,
        "line": {"color": COLORS["Essential System Services"], "width": 0.5},
        "fillcolor": COLORS["Essential System Services"], "stackgroup": "dispatch", "stackgaps": "infer zero",
        "customdata": list(map(list, zip(*(values(col) for col in reserve_columns)))),
        "meta": {"sourceColumns": reserve_columns, "sourceValueCount": 97, "derivedExpression": "U + V + Z + AA"},
        "hovertemplate": "%{x|%d %b %Y %H:%M}<br>Essential System Services: %{y:,.2f} MW<extra></extra>",
    })
    spec["data"] += [trace("N", "Thermal headroom", stack="dispatch"), trace("O", "Forced outages", stack="dispatch"),
                     trace("R", "ESR withdrawal", color="#AD56A7", negate=True, fill=True),
                     trace("P", color="#111111"), trace("Q", color="#111111", dash="dashdot"),
                     trace("S", "ESR state of charge", color="#AD56A7", dash="dash", secondary=True)]
    spec["layout"] = layout("Generation, storage and demand (MW)")
    spec["layout"]["yaxis"]["range"] = [-2000, 8000]
    spec["layout"]["yaxis2"] = {"title": {"text": "State of charge (%)"}, "overlaying": "y", "side": "right",
                                "range": [-25, 100], "tickvals": [0, 25, 50, 75, 100], "showgrid": False}
    spec["layout"]["margin"]["b"] = 210
    spec["layout"]["legend"]["y"] = -0.3
    spec["layout"]["height"] = 720
    spec["sources"] = base_sources + [source(col) for col in list("DEFGHIJKLMNOPQRS") + reserve_columns]
    spec["notes"] = common_notes + [
        "The Essential System Services area is the sum of gas and ESR Regulation Raise and Contingency Reserve Raise (columns U, V, Z and AA), reproducing the source figure's reserve layer.",
    ]

    battery = base_spec("supplement-a4-f28-battery", sheet.title,
                        "Figure 28 — Battery utilisation during the two-day unserved-energy event", 37)
    battery["subtitle"] = "Middle panel · discharge, reserves, charging and state of charge"
    battery["data"] = [trace("J", "Injection", stack="battery", color="#AE68AB"),
                       trace("Z", "Regulation Raise", stack="battery"),
                       trace("AA", "Contingency Reserve Raise", stack="battery"),
                       trace("Y", "Planned Outage", stack="battery"),
                       trace("X", "Forced Outage", stack="battery"),
                       trace("R", "Withdrawal", color="#AD56A7", negate=True, fill=True),
                       trace("AB", "Available capacity", color="#111111"),
                       trace("S", "ESR state of charge", color="#AD56A7", dash="dash", secondary=True)]
    battery["layout"] = layout("Battery power (MW)")
    battery["layout"]["yaxis"]["range"] = [-1200, 1500]
    battery["layout"]["yaxis2"] = {"title": {"text": "State of charge (%)"}, "overlaying": "y", "side": "right",
                                   "range": [0, 100], "tickmode": "array", "tickvals": [0, 20, 40, 60, 80, 100],
                                   "showgrid": False}
    battery["sources"] = base_sources + [source(col) for col in ["J", "Z", "AA", "Y", "X", "R", "AB", "S"]]
    battery["notes"] = common_notes.copy()

    gas = base_spec("supplement-a4-f28-gas", sheet.title,
                    "Figure 28 — Gas utilisation during the two-day unserved-energy event", 37)
    gas["subtitle"] = "Bottom panel · generation, reserves, headroom and outages"
    gas["data"] = [trace("H", "Generation", stack="gas", color="#59C8B7"),
                   trace("U", "Regulation Raise", stack="gas"),
                   trace("V", "Contingency Reserve Raise", stack="gas"),
                   trace("T", "Headroom", stack="gas"),
                   trace("O", "Forced Outage", stack="gas"),
                   trace("W", "Available capacity", color="#111111")]
    gas["layout"] = layout("Gas power (MW)")
    gas["layout"]["yaxis"]["range"] = [0, 3200]
    gas["sources"] = base_sources + [source(col) for col in ["H", "U", "V", "T", "O", "W"]]
    gas["notes"] = [common_notes[0],
                    "The worksheet's shared Outages series is used for forced outages. At all 97 samples, it exactly equals the difference between the maximum gas capacity and Gas Available Capacity. No separate nonzero gas planned-outage series is supplied."]
    # Verify the three independent source identities before producing output.
    gas_max = max(values("W"))
    assert all(math.isclose(o + w, gas_max, abs_tol=1e-7) for o, w in zip(values("O"), values("W")))
    gas_balance = [h + u + v + t - w for h, u, v, t, w in zip(*(values(c) for c in ["H", "U", "V", "T", "W"]))]
    battery_balance = [j + z + a - ab for j, z, a, ab in zip(*(values(c) for c in ["J", "Z", "AA", "AB"]))]
    demand_balance = [p + r - q for p, r, q in zip(*(values(c) for c in ["P", "R", "Q"]))]
    for errors in [gas_balance, battery_balance, demand_balance]:
        assert max(map(abs, errors)) < 1e-7
    return [spec, battery, gas], {
        "half_hour_samples": len(rows), "start": times[0], "end": times[-1],
        "gas_available_capacity_identity_max_error_mw": max(map(abs, gas_balance)),
        "battery_available_capacity_identity_max_error_mw": max(map(abs, battery_balance)),
        "operational_demand_identity_max_error_mw": max(map(abs, demand_balance)),
        "gas_max_capacity_mw": gas_max,
    }


def main() -> None:
    workbook = load_workbook(BOOK, read_only=True, data_only=True)
    bubble = bubble_spec(workbook)
    temporal, balance_review = temporal_specs(workbook)
    specs = [bubble] + temporal
    target = ROOT / "src/charts/supplemental_duration.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(specs, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    review = {
        "worksheet_figures": ["A4_F.27", "A4_F.28"], "interactive_specs": len(specs),
        "pdf_pages_visually_reviewed": [62, 63], "workbook_images_visually_reviewed": [36, 37],
        "bubble_observations": 300, "positive_frequency_observations": 39, "zero_frequency_placeholders": 261,
        "frequency_total": 1.0, "no_event_frequency": 0.711875,
        "temporal_source_identity_checks": balance_review,
        "derived_elements": [
            "Reliability shading: magnitude (%) × duration (hours) = 8760 × 0.0002 = 1.752.",
            "Dispatch Essential System Services: Gas Regulation Raise + Gas Contingency Reserve Raise + ESR Regulation Raise + ESR Contingency Reserve Raise.",
            "ESR withdrawal signs inverted for display; original positive values retained as customdata.",
        ],
    }
    (ROOT / "data/analysis/supplemental-duration-review.json").write_text(json.dumps(review, indent=2) + "\n")
    print(f"Built {len(specs)} supplemental charts in {target.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
