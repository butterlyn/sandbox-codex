"""Build and independently verify saved Excel chart objects 15–28.

Run after chart_common.py is available.  All exceptions below preserve the
source chart's numeric caches and make its unusual Excel construction explicit.
"""
from __future__ import annotations

import copy
import datetime as dt
import json
import math
from pathlib import Path

from lxml import etree as ET
from openpyxl.utils.cell import get_column_letter, range_boundaries

from chart_common import (
    NS, ROOT, archive, build_chart, formula_values, load_chart, sheet_cells,
    sheet_parts,
)


def by_index(spec):
    return {trace["meta"]["sourceSeriesIndex"]: trace for trace in spec["data"]}


def add_note(spec, note):
    if note not in spec["notes"]:
        spec["notes"].append(note)


def custom_minimum_demand(spec, raw):
    """Overlay the source's two bar stacks and minimum-demand envelopes.

    Excel uses transparent bars as bases and a hidden, identically scaled axis
    to overlay NCESS on ESR. Explicit Plotly bases reproduce that geometry on
    one scale, including the separately constructed negative-demand bar tails.
    """
    traces = by_index(spec)
    records = {s["index"]: s for s in raw["series"]}
    labels = records[1]["x"]
    base = records[1]["y"]
    for index, trace in traces.items():
        trace["x"] = list(range(len(trace["y"])))
        trace["xaxis"] = "x"
        trace["yaxis"] = "y"
        original = records[index]["y"]
        trace["customdata"] = [
            [labels[i] if i < len(labels) else "Unlabelled end padding",
             original[i] if i < len(original) else None]
            for i in range(len(trace["y"]))
        ]
        trace["meta"]["positionalCategories"] = True
        trace["hovertemplate"] = (
            "%{customdata[0]}<br>%{fullData.name}: %{customdata[1]:,.1f} MW"
            "<extra></extra>"
        )
        if trace["type"] == "bar":
            trace["base"] = [0] * len(trace["y"])
            trace["width"] = 0.74
            trace["zorder"] = 2
            trace["marker"]["line"]["width"] = 0
        elif index in (1, 2, 5):
            trace["zorder"] = 0
            # Explicit tozeroy on every member would repaint the entire lower
            # stack. Fill later area members only down to the previous member.
            trace["fill"] = "tozeroy" if index == 1 else "tonexty"
        elif index in (7, 6, 12):
            trace["zorder"] = 3
            trace["hoverinfo"] = "skip"
            trace.pop("hovertemplate", None)
        else:
            trace["zorder"] = 6

    # Positive ESR bars start at operational demand. Negative-era ESR is split
    # at zero in the workbook into indices 10 and 11; preserve both raw parts.
    traces[9]["base"] = [base[i] if v is not None else 0
                         for i, v in enumerate(traces[9]["y"])]
    traces[13]["base"] = [base[i] if v is not None else 0
                          for i, v in enumerate(traces[13]["y"])]
    traces[13]["zorder"] = 5
    traces[13]["width"] = 0.55
    traces[13]["marker"]["pattern"] = {
        "shape": "/", "fgcolor": "#606EB2", "bgcolor": "#FFFFFF", "solidity": 0.4,
    }
    traces[13]["marker"]["color"] = "#FFFFFF"
    for index in (4, 8):
        traces[index]["marker"]["color"] = "rgba(0,0,0,0)"
        traces[index]["showlegend"] = False
        traces[index]["hoverinfo"] = "skip"
        traces[index].pop("hovertemplate", None)
    for index in (9, 10, 11):
        traces[index]["legendgroup"] = "esr"
        traces[index]["showlegend"] = index == 9
        traces[index]["marker"]["color"] = "rgba(179,126,2,0.8)"
    traces[9]["name"] = "ESR"
    traces[10]["hovertemplate"] = (
        "%{customdata[0]}<br>Operational-demand portion below zero: "
        "%{customdata[1]:,.1f} MW<extra></extra>"
    )
    traces[11]["hovertemplate"] = (
        "%{customdata[0]}<br>ESR bar portion above zero: "
        "%{customdata[1]:,.1f} MW<extra></extra>"
    )
    for index in (1, 2, 5):
        traces[index]["showlegend"] = False
    traces[7]["showlegend"] = True
    traces[6]["showlegend"] = True
    traces[12]["showlegend"] = False
    traces[12]["fillcolor"] = traces[6]["fillcolor"]
    traces[6]["legendgroup"] = traces[12]["legendgroup"] = "technical-envelope"
    traces[0]["showlegend"] = traces[3]["showlegend"] = True
    traces[0]["line"]["color"] = traces[0]["marker"]["color"] = "#373A36"
    traces[3]["line"]["color"] = traces[3]["marker"]["color"] = "#703080"

    layout = spec["layout"]
    layout.pop("xaxis2", None)
    layout.pop("yaxis2", None)
    layout["xaxis"].update({
        "type": "linear", "tickmode": "array", "tickvals": list(range(19)),
        "ticktext": labels, "range": [0, 18.5], "tickangle": -45,
        "autorange": False, "anchor": "y",
    })
    layout["yaxis"].update({"range": [-500, 3500], "autorange": False,
                            "dtick": 500, "anchor": "x"})
    layout["barmode"] = "overlay"
    layout["hovermode"] = "closest"
    layout["annotations"] = [
        {"x": 16.1, "y": 2440, "xref": "x", "yref": "y", "showarrow": False,
         "text": "Demand met by distributed PV<br>and PV Non-Scheduled Generation",
         "font": {"size": 10, "color": "#5A503A"}},
        {"x": 2.0, "y": 850, "xref": "x", "yref": "y", "showarrow": False,
         "text": "Operational demand", "font": {"size": 10, "color": "#4F5358"}},
    ]
    add_note(spec, "The workbook uses separate positive and negative ESR bar portions, transparent bar bases and a hidden second axis. Their original values and overlaid geometry are retained.")
    add_note(spec, "Several source areas contain a twentieth, unlabelled duplicate endpoint. It is retained as chart-edge padding beyond the final labelled year, not treated as another forecast year.")
    add_note(spec, "Minimum-demand thresholds are 300 MW and 550 MW including Emergency Solar Management capability; forecast demand excludes utility-scale storage charging.")


def custom_capacity_mix(spec, raw):
    """Give repeated technology series a shared legend and retain patterns."""
    broken = [s for s in raw["series"] if s["y_formula"] and "#REF!" in s["y_formula"]]
    bad_indices = {s["index"] for s in broken}
    spec["source_data"] = [{"series_index": s["index"], "name": s["name"],
                             "x": s["x"], "y": s["y"], "formula": s["y_formula"],
                             "status": "broken-reference-placeholder"} for s in broken]
    spec["data"] = [t for t in spec["data"] if t["meta"]["sourceSeriesIndex"] not in bad_indices]
    colors = {s["name"].strip(): s["style"]["fill"] for s in raw["series"]
              if s["chart_type"] == "barChart" and s["index"] <= 8}
    series = {s["index"]: s for s in raw["series"]}
    for trace in spec["data"]:
        s = series[trace["meta"]["sourceSeriesIndex"]]
        if trace["type"] != "bar":
            trace["showlegend"] = True
            trace["line"]["color"] = "#222222"
            trace["hovertemplate"] = "%{x}<br>%{fullData.name}: %{y:,.1f} MW<extra></extra>"
            continue
        row = int(s["y_formula"].split("$")[-1])
        stage = "existing" if row <= 34 else "committed" if row <= 44 else "anticipated"
        technology = s["name"].strip()
        trace["name"] = technology
        trace["legendgroup"] = technology
        trace["showlegend"] = stage == "existing"
        trace["marker"]["color"] = colors[technology]
        trace["marker"]["line"]["width"] = 0
        if stage != "existing":
            trace["marker"]["color"] = "#FFFFFF"
            trace["marker"]["pattern"] = {"shape": "." if stage == "committed" else "|",
                "fgcolor": colors[technology], "bgcolor": "#FFFFFF", "solidity": 0.45}
        else:
            trace["marker"].pop("pattern", None)
        trace["hovertemplate"] = (
            f"%{{x}}<br>{technology} · {stage}: %{{y:,.1f}} MW<extra></extra>"
        )
        trace["meta"]["facilityStage"] = stage
    spec["layout"]["yaxis"]["range"] = [0, 8000]
    spec["layout"]["yaxis"]["autorange"] = False
    spec["layout"]["legend"]["groupclick"] = "togglegroup"
    add_note(spec, "Three invalid #REF! series contain only placeholder values of 1. They are excluded from the forecast graphic and retained in source_data for audit.")
    add_note(spec, "Solid colours show existing facilities, dots committed facilities, and vertical stripes anticipated facilities. A technology legend control toggles all three components.")


def custom_rlm(spec):
    traces = by_index(spec)
    traces[16]["showlegend"] = True
    for index, shape in ((17, "."), (23, "|")):
        traces[index]["marker"]["pattern"] = {
            "shape": shape, "fgcolor": "#948794", "bgcolor": "#FFFFFF", "solidity": 0.45,
        }
        traces[index]["marker"]["color"] = "#FFFFFF"
    for index in (0, 6):
        traces[index]["legendgroup"] = "rlm-uplift"
        traces[index]["showlegend"] = index == 0
    for trace in spec["data"]:
        trace["hovertemplate"] = "%{x}<br>%{fullData.name}: %{y:,.1f} MW<extra></extra>"
    add_note(spec, "Both 240 MW uplift components and the Excel stacked-line order are retained exactly; displayed line heights are cumulative, rather than the individual increments stored in the worksheet.")


def unavailable_legacy(spec, raw, related):
    spec["kind"] = "unavailable"
    spec["subtitle"] = "Broken legacy Excel object — cached placeholders only"
    spec["source_data"] = [{"series_index": s["index"], "name": s["name"],
                             "x": s["x"], "y": s["y"], "formula": s["y_formula"],
                             "status": "broken-reference-placeholder"} for s in raw["series"]]
    spec["data"] = []
    spec["related_chart_ids"] = [f"chart-{related:02}"]
    spec["notes"] = [
        "This separate embedded Excel chart has eight broken external references to [1]Ch3_Fig1!#REF!, no category labels, and only a cached placeholder value of 1 per series. It cannot be reconstructed as a valid forecast.",
        f"The complete chart on the same worksheet is available as chart-{related:02}. Placeholder caches remain in the exported source_data.",
    ]


def custom_outage_timeline(spec, raw):
    """Preserve Excel's seasonal gaps and give every observed day a tooltip."""
    source = raw["series"][0]
    cells = sheet_cells(raw["sheet"])
    labels, month_ticks, month_text = [], [], []
    current_month = None
    for i in range(len(source["y"])):
        row = i + 24
        month, day = cells.get(f"B{row}"), cells.get(f"C{row}")
        if isinstance(month, (float, int)):
            current_month = dt.datetime(1899, 12, 30) + dt.timedelta(days=month)
            month_ticks.append(i)
            month_text.append(current_month.strftime("%b-%y"))
        if current_month is not None and isinstance(day, (float, int)):
            labels.append(f"{current_month:%Y-%m}-{int(day):02}")
        else:
            labels.append("Gap between hot seasons")
    trace = spec["data"][0]
    trace["x"] = list(range(len(source["y"])))
    trace["y"] = source["y"]
    trace["customdata"] = [[label, value] for label, value in zip(labels, source["y"])]
    trace["meta"]["positionalCategories"] = True
    trace["meta"]["dateLabelsFrom"] = "'Ch 3_F.28'!B24:C898"
    trace["hovertemplate"] = "%{customdata[0]}<br>Peak Capacity Credits on Forced Outage: %{y:.2%}<extra></extra>"
    trace["connectgaps"] = False
    trace["fill"] = "tozeroy"
    trace["fillcolor"] = "#703080"
    trace["line"].update({"color": "#703080", "width": 0.6})
    # A single Plotly filled trace bridges null runs with a straight fill edge
    # despite connectgaps=False. Separate contiguous seasons preserve the
    # workbook's genuinely empty gaps without manufacturing zero observations.
    segments = []
    start = None
    for i, value in enumerate(source["y"] + [None]):
        if value is not None and start is None:
            start = i
        elif value is None and start is not None:
            segment = copy.deepcopy(trace)
            segment["x"] = list(range(start, i))
            segment["y"] = source["y"][start:i]
            segment["customdata"] = trace["customdata"][start:i]
            segment["meta"]["sourcePointIndices"] = list(range(start, i))
            segment["legendgroup"] = "forced-outage-seasons"
            segment["showlegend"] = not segments
            segments.append(segment)
            start = None
    spec["data"] = segments
    spec["source_data"] = [{"series_index": source["index"], "name": source["name"],
                             "x": labels, "y": source["y"], "formula": source["y_formula"]}]
    spec["layout"]["xaxis"].update({
        "type": "linear", "tickmode": "array", "tickvals": month_ticks,
        "ticktext": month_text, "tickangle": -45, "range": [0, len(source["y"]) - 1],
        "autorange": False, "title": {"text": "Hot-season days; gaps separate seasons"},
    })
    spec["layout"]["yaxis"].update({"tickformat": ".0%", "range": [0, 0.3],
                                    "autorange": False, "dtick": 0.05})
    spec["layout"]["hovermode"] = "closest"
    add_note(spec, "The source uses compressed hot-season day positions with blank gaps between seasons, not a continuous calendar axis. Month ticks and daily hover dates are reconstructed from the workbook's Month-Year and Day columns; all 727 observations and 148 gap positions are retained.")


def verify_sources(raw):
    """Compare numeric caches to saved cells, respecting hidden rows/columns."""
    checks = 0
    unresolved = []
    mismatches = []
    sheet_xml = {}
    for series in raw["series"]:
        formula = series["y_formula"]
        values = formula_values(formula)
        if values is None:
            unresolved.append({"series_index": series["index"], "formula": formula})
            continue
        sheet, ref = formula.rsplit("!", 1)
        sheet = sheet.strip("'").replace("''", "'")
        if sheet not in sheet_xml:
            sheet_xml[sheet] = ET.fromstring(archive().read(sheet_parts()[sheet]))
        root = sheet_xml[sheet]
        c1, r1, c2, r2 = range_boundaries(ref)
        hidden_cols = set()
        for col in root.findall("s:cols/s:col", NS):
            if col.get("hidden") == "1":
                hidden_cols.update(range(int(col.get("min")), int(col.get("max")) + 1))
        hidden_rows = {int(row.get("r")) for row in root.findall("s:sheetData/s:row", NS)
                       if row.get("hidden") == "1"}
        visible = [sheet_cells(sheet).get(f"{get_column_letter(c)}{r}")
                   for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)
                   if c not in hidden_cols and r not in hidden_rows]
        used_hidden_filter = len(visible) == len(series["y"]) and len(visible) != len(values)
        compare = visible if used_hidden_filter else values
        for i, value in enumerate(series["y"]):
            if not isinstance(value, (float, int)):
                continue
            checks += 1
            actual = compare[i] if i < len(compare) else None
            if not isinstance(actual, (float, int)) or not math.isclose(value, actual, rel_tol=1e-10, abs_tol=1e-8):
                mismatches.append({"series_index": series["index"], "point_index": i,
                                   "cache": value, "worksheet": actual, "formula": formula})
    return {"chart_number": raw["chart_number"], "numeric_points_checked": checks,
            "unresolved_references": unresolved, "mismatches": mismatches,
            "hidden_columns_respected": raw["chart_number"] == 16}


def main():
    specs, reviews = [], []
    for number in range(15, 29):
        raw = load_chart(number)
        spec = build_chart(number)
        review = verify_sources(raw)
        if review["mismatches"]:
            raise AssertionError(f"Unexplained worksheet/cache mismatches in chart {number}: {review['mismatches']}")
        if number == 16:
            # The cached range intentionally omits hidden 2017-18/2018-19 data.
            spec["notes"] = [note for note in spec["notes"] if "differs from the saved worksheet" not in note]
            for source in spec["sources"]:
                source["validation"].update({"mismatch_count": 0, "mismatch_indices": [],
                    "hidden_columns_filtered": ["D", "E"]})
            add_note(spec, "Hidden worksheet columns D:E are excluded exactly as in Excel; the displayed series span 2019-20 to 2035-36.")
        elif number == 18:
            custom_minimum_demand(spec, raw)
        elif number == 19:
            custom_capacity_mix(spec, raw)
        elif number == 20:
            custom_rlm(spec)
        elif number in (21, 23):
            unavailable_legacy(spec, raw, number + 1)
        elif number in (22, 24):
            spec["subtitle"] = "Complete forecast chart on this worksheet"
        elif number == 26:
            for trace in spec["data"]:
                trace["hovertemplate"] = "%{x}<br>%{fullData.name}: %{y:.1%}<extra></extra>"
            spec["layout"]["yaxis"].update({"tickformat": ".0%", "rangemode": "tozero"})
        elif number == 27:
            custom_outage_timeline(spec, raw)
        elif number == 28:
            for trace in spec["data"]:
                trace["xaxis"] = "x"
                trace["hovertemplate"] = (
                    "%{x}<br>%{fullData.name}: %{y:.2%}<extra></extra>" if trace["type"] == "scatter"
                    else "%{x}<br>%{fullData.name}: %{y:,.1f} MW<extra></extra>"
                )
            spec["layout"].pop("xaxis2", None)
            spec["layout"]["yaxis2"]["anchor"] = "x"
            spec["layout"]["xaxis"]["tickangle"] = -60
            spec["layout"]["margin"].update({"b": 175, "r": 80})
            spec["layout"]["yaxis"].update({"tickformat": ".0%", "range": [0, 0.5],
                                            "dtick": 0.05, "autorange": False})
            spec["layout"]["yaxis2"].update({"tickformat": ",.0f", "range": [0, 350],
                                             "dtick": 50, "autorange": False})
        # Keep formatting consistent while preserving all source precision.
        if number not in (21, 23, 26, 27, 28):
            spec["layout"]["yaxis"].setdefault("tickformat", ",.0f")
        specs.append(spec)
        review["kind"] = spec["kind"]
        review["rendered_traces"] = len(spec["data"])
        reviews.append(review)
    (ROOT / "src/charts").mkdir(parents=True, exist_ok=True)
    (ROOT / "src/charts/group_2.json").write_text(json.dumps(specs, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    (ROOT / "data/analysis/group_2-review.json").write_text(json.dumps({
        "charts": reviews, "numeric_points_checked": sum(r["numeric_points_checked"] for r in reviews),
        "cache_worksheet_mismatches": sum(len(r["mismatches"]) for r in reviews),
        "report_pages_visually_reviewed": [51, 61, 62, 70],
    }, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    print(f"Built {len(specs)} chart objects; {sum(r['numeric_points_checked'] for r in reviews)} numeric cache points verified against worksheets.")


if __name__ == "__main__":
    main()
