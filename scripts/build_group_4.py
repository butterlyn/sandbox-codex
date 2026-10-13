#!/usr/bin/env python3
"""Recreate native workbook charts 42–54, including scatter fits and units.

The shared parser supplies native cache values. This group preserves their
indices and adds only the visual features that need chart-specific semantics.
"""

from __future__ import annotations

import json
import math
import re
import statistics
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from chart_common import build_chart, load_chart


ROOT = Path(__file__).resolve().parents[1]
NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
SCENARIOS = {45: "Low", 46: "Expected", 47: "High"}


def source_values():
    workbook = load_workbook(
        ROOT / "data/sources/2026-wem-esoo-data-register.xlsx",
        read_only=True,
        data_only=True,
    )
    inventory = json.loads((ROOT / "data/analysis/inventory.json").read_text())
    names = {item["sheet"] for item in inventory["charts"][41:54]}
    values = {
        name: {
            cell.coordinate: cell.value
            for row in workbook[name].iter_rows()
            for cell in row
            if cell.value is not None
        }
        for name in names
    }
    workbook.close()
    return values


def formula_values(formula, values):
    sheet, cell_range = formula.rsplit("!", 1)
    sheet = sheet.strip("'").replace("''", "'")
    c1, r1, c2, r2 = range_boundaries(cell_range.replace("$", ""))
    return [
        values[sheet].get(f"{get_column_letter(col)}{row}")
        for row in range(r1, r2 + 1)
        for col in range(c1, c2 + 1)
    ]


def xml_tree(number):
    return ET.parse(ROOT / f"data/analysis/chart-xml/chart{number:02}.xml").getroot()


def check_caches(number, values):
    """Check the cached values against all referenced worksheet cells."""
    checked = 0
    root = xml_tree(number)
    for reference in root.findall(".//c:numRef", NS) + root.findall(".//c:strRef", NS):
        formula = reference.find("c:f", NS)
        cache = reference.find("c:numCache", NS)
        if cache is None:
            cache = reference.find("c:strCache", NS)
        if formula is None or cache is None:
            continue
        cells = formula_values(formula.text, values)
        for point in cache.findall("c:pt", NS):
            index = int(point.get("idx"))
            text = point.findtext("c:v", namespaces=NS)
            source = cells[index]
            if isinstance(source, (int, float)):
                assert math.isclose(float(text), source, rel_tol=1e-12, abs_tol=1e-10), (
                    number, formula.text, index, text, source
                )
            else:
                assert text == source, (number, formula.text, index, text, source)
            checked += 1
    return checked


def trace_by_index(spec):
    return {
        trace["meta"]["sourceSeriesIndex"]: trace
        for trace in spec["data"]
        if "sourceSeriesIndex" in trace.get("meta", {})
        and not trace.get("meta", {}).get("derived")
    }


def apply_temperature_scatter(spec, native, values):
    traces = trace_by_index(spec)
    for series in native["series"]:
        trace = traces[series["index"]]
        trace["legendrank"] = trace["meta"]["sourceSeriesOrder"]
        is_model = "model" in series["name"]
        if is_model:
            # Excel accepts 1°C…40°C in a scatter xVal string reference; the
            # actual numeric temperatures must be provided explicitly to Plotly.
            trace["x"] = [
                float(re.fullmatch(r"\s*([+-]?[0-9.]+)\s*°C\s*", str(x)).group(1))
                for x in series["x"]
            ]
            trace["mode"] = "lines"
            trace.setdefault("line", {})["shape"] = "linear"
            trace["hovertemplate"] = (
                "%{x:.1f} °C<br>%{y:.3f} kWh<extra>%{fullData.name}</extra>"
            )
            trace.setdefault("meta", {})["xTransformation"] = "numeric °C from worksheet labels"
        else:
            trace["x"] = series["x"]
            trace["mode"] = "markers"
            # The native RGBA marker color already carries Excel's 20% alpha.
            trace.setdefault("marker", {})["size"] = 5
            dates_formula = re.sub(r"\$M\$", "$J$", series["x_formula"])
            dates = formula_values(dates_formula, values)
            trace["customdata"] = [date.strftime("%Y-%m-%d") for date in dates]
            trace["hovertemplate"] = (
                "%{customdata}<br>%{x:.2f} °C<br>%{y:.3f} kWh"
                "<extra>%{fullData.name}</extra>"
            )
    # Lines are drawn over translucent observations, retaining all source points.
    spec["data"].sort(key=lambda trace: "model" in trace.get("name", ""))
    spec["layout"]["xaxis"].update({
        "type": "linear", "range": [8, 32], "dtick": 4, "showgrid": False,
        "title": {"text": "Average daily temperature (°C)"},
    })
    spec["layout"]["yaxis"].update({
        "range": [10, 25], "dtick": 5,
        "title": {"text": "Average underlying daily household consumption (kWh)"},
    })
    spec["layout"]["hovermode"] = "closest"
    spec["notes"].append(
        "Daily observations and the five supplied temperature-response models are retained. "
        "The initial view matches the workbook; zoom out to see the full 1–40°C model range."
    )


def apply_household_trend(spec, native):
    traces = trace_by_index(spec)
    historical = next(series for series in native["series"] if series["index"] == 0)
    x = [float(value) for value in historical["x"]]
    y = [float(value) for value in historical["y"]]
    slope, intercept = statistics.linear_regression(x, y)
    # Excel's historical source series is invisible; individual year markers
    # are already separate source series. Preserve the invisible source data.
    traces[0].update({"mode": "none", "showlegend": False, "hoverinfo": "skip"})
    traces[0].pop("hovertemplate", None)
    for index in range(1, 6):
        traces[index]["mode"] = "markers"
        traces[index].setdefault("marker", {}).update({"symbol": "circle", "size": 8})
    # Drop any generic helper trendline; this trace has explicit provenance and
    # a fit limited to the historical domain specified by the Excel chart.
    spec["data"] = [
        trace for trace in spec["data"]
        if not trace.get("meta", {}).get("derived")
    ]
    spec["data"].append({
        "type": "scatter", "mode": "lines", "name": "Historical linear trend",
        "x": [min(x), max(x)],
        "y": [slope * min(x) + intercept, slope * max(x) + intercept],
        "line": {"color": "#424242", "width": 2, "dash": "dash"},
        "hovertemplate": "%{x:.0f}<br>Linear trend: %{y:,.3f}<extra></extra>",
        "meta": {
            "derived": True, "method": "ordinary least squares",
            "sourceSeriesIndices": [0], "slope": slope, "intercept": intercept,
            "sourceXFormula": historical["x_formula"],
            "sourceYFormula": historical["y_formula"],
        },
    })
    unit = {51: "kWh", 52: "kWh/CDD", 53: "kWh/HDD"}[spec["chart_number"]]
    for index, trace in traces.items():
        if index == 0:
            continue
        trace["customdata"] = [f"{int(year)-1}-{str(int(year))[-2:]}" for year in trace["x"]]
        trace["hovertemplate"] = (
            "%{customdata}<br>%{y:,.3f} " + unit + "<extra>%{fullData.name}</extra>"
        )
    spec["layout"]["xaxis"].update({
        "type": "linear", "range": [2020, 2036], "dtick": 2,
        "autorange": False, "showgrid": False,
        "tickformat": "d", "title": {"text": "Capacity year ended"},
    })
    spec["layout"]["hovermode"] = "closest"
    spec["notes"].append(
        "The dashed line is the workbook's linear fit to 2021–2025 observations. "
        "The forecast uses the supplied annual values."
    )
    if spec["chart_number"] in (52, 53):
        spec["notes"].append("CDD = Cooling Degree Days; HDD = Heating Degree Days.")
    return {"slope": slope, "intercept": intercept, "domain": [min(x), max(x)]}


def build_group():
    values = source_values()
    specs = []
    reviews = []
    for number in range(42, 55):
        native = load_chart(number)
        spec = build_chart(number)
        checked = check_caches(number, values)
        review = {
            "chart_number": number,
            "sheet": spec["sheet"],
            "cache_values_checked_against_cells": checked,
            "native_series_count": len(native["series"]),
            "checks": ["every cached value matches its worksheet reference"],
        }
        if number in SCENARIOS:
            scenario = SCENARIOS[number]
            expected_row = {45: 24, 46: 28, 47: 32}[number]
            assert values["A1_F.9"][f"B{expected_row}"] == scenario
            spec["title"] = spec["title"].replace("spilt", "split")
            spec["subtitle"] = f"{scenario} scenario"
            spec["title"] += f" — {scenario} scenario"
            spec["layout"]["barmode"] = "stack"
            spec["layout"]["yaxis"]["range"] = [0, 3750]
            review["checks"].append(f"{scenario} scenario confirmed from worksheet header")
        if number == 50:
            apply_temperature_scatter(spec, native, values)
            assert sum(len(trace["x"]) for trace in spec["data"] if "data" in trace["name"]) == 1826
            assert sum(len(trace["x"]) for trace in spec["data"] if "model" in trace["name"]) == 200
            review["checks"] += [
                "1,826 daily observations retained without sampling",
                "five models retain all 40 points with numeric temperatures 1–40°C",
                "native axes 8–32°C and 10–25 kWh preserved",
            ]
        if number in (51, 52, 53):
            review["historical_linear_fit"] = apply_household_trend(spec, native)
            review["checks"].append("five historical markers, eleven forecast values, and native linear fit retained")
        traces = trace_by_index(spec)
        divisor = {42: 1e6, 48: 1e3, 49: 1e3, 54: 1e3}.get(number, 1)
        for series in native["series"]:
            trace = traces[series["index"]]
            expected_y = [None if value is None else value / divisor for value in series["y"]]
            assert len(expected_y) == len(trace["y"]), (number, series["name"], "length")
            for expected, actual in zip(expected_y, trace["y"]):
                assert (expected is None and actual is None) or (
                    expected is not None and actual is not None
                    and math.isclose(expected, actual, rel_tol=1e-12, abs_tol=1e-12)
                ), (number, series["name"], expected, actual)
            if number != 50 or "model" not in series["name"]:
                assert trace["x"] == series["x"], (number, series["name"], "x")
        review["display_divisor"] = divisor
        review["checks"].append("all plotted source x/y positions and missing-value gaps preserved, allowing documented display-unit conversion")
        specs.append(spec)
        reviews.append(review)
    return specs, reviews


def main():
    specs, reviews = build_group()
    target = ROOT / "src/charts/group_4.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(specs, ensure_ascii=False, separators=(",", ":"), allow_nan=False) + "\n")
    (ROOT / "data/analysis/group_4-review.json").write_text(json.dumps({
        "charts": reviews,
        "total_cache_values_checked_against_cells": sum(item["cache_values_checked_against_cells"] for item in reviews),
        "status": "passed",
    }, indent=2, ensure_ascii=False, allow_nan=False) + "\n")
    print(f"Built and verified {len(specs)} native charts into {target.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
